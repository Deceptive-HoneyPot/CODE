"""
Builds paper.docx (single column) and paper_ieee.docx (two column) from the
generated paper.tex. paper.tex/paper.pdf remain canonical; the .docx files
are editable approximations: equations and algorithms are rendered as text,
tables as Word tables, figures embedded from figures/.

    /usr/bin/python3 build_docx.py     (after build_latex.py)
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).parent
TEX = (ROOT / "paper.tex").read_text()

MATH = {r"\ge": "≥", r"\le": "≤", r"\ne": "≠", r"\in": "∈", r"\cup": "∪", r"\wedge": "∧", r"\Leftrightarrow": "⇔",
        r"\gets": "←", r"\to": "→", r"\times": "×", r"\pm": "±", r"\cdot": "·", r"\emptyset": "∅",
        r"\lceil": "⌈", r"\rceil": "⌉", r"\langle": "⟨", r"\rangle": "⟩", r"\|": "‖", r"\min": "min",
        r"\mathcal{L}": "L", r"\Ledger": "L", r"\Obs": "Ω", r"\Omega": "Ω", r"\kappa": "κ", r"\tau": "τ",
        r"\rho": "ρ", r"\alpha": "α", r"\beta": "β", r"\sigma": "σ", r"\sqrt": "√", r"\bullet": "●",
        r"\mu": "μ", r"\mathcal{P}": "P", r"\mathcal{R}": "R", r"\dots": "…", r"\,": " ", r"\;": " ", r"\quad": "  ",
        r"\qquad": "    ", r"\textbf": "", r"\textit": "", r"\emph": "", r"\mathrm": "", r"\mathit": "", r"\text": "",
        r"\big": "", r"\left": "", r"\right": "", r"\log_2": "log2", r"\{": "{", r"\}": "}", r"\%": "%", r"\&": "&",
        r"\_": "_", r"\#": "#", r"\shortstack": "", r"\ldots": "…"}


def detex(s: str) -> str:
    s = re.sub(r"~?\\cite\{([^}]*)\}", lambda m: "[" + ", ".join(k.replace("ref", "") for k in m.group(1).split(",")) + "]", s)
    s = re.sub(r"(~?)\\ref\{([^}]*)\}", lambda m: (" " if m.group(1) else "") + REFS.get(m.group(2), "?"), s)
    s = re.sub(r"\\label\{[^}]*\}", "", s)
    s = re.sub(r"\\(texttt|textbf|textit|emph|mathrm|mathit|text|url|textsuperscript)\{([^{}]*)\}", r"\2", s)
    s = re.sub(r"\\(texttt|textbf|textit|emph|mathrm|mathit|text|url)\{([^{}]*)\}", r"\2", s)
    for k in sorted(MATH, key=len, reverse=True):
        s = s.replace(k, MATH[k])
    s = s.replace("$", "").replace("--", "–").replace("~", " ").replace("\\\\", " ")
    s = re.sub(r"\\[a-zA-Z]+", "", s)
    s = re.sub(r"[{}]", "", s)
    return re.sub(r"[ \t]+", " ", s).strip()


# numbering for \ref
REFS = {}
for i, m in enumerate(re.finditer(r"\\section\{([^}]*)\}\s*\\label\{([^}]*)\}", TEX), 1):
    pass
sec_no = 0; sub_no = 0; fig_no = 0; tab_no = 0; alg_no = 0; eq_no = 0
roman = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX"]
for m in re.finditer(r"\\(section|subsection)\{([^}]*)\}(?:\s*\\label\{([^}]*)\})?|\\begin\{(figure\*?|table\*?|algorithm)\}.*?\\label\{([^}]*)\}|\\begin\{equation\}.*?\\label\{([^}]*)\}", TEX, re.S):
    if m.group(1) == "section":
        sec_no += 1; sub_no = 0
        if m.group(3): REFS[m.group(3)] = roman[sec_no - 1]
    elif m.group(1) == "subsection":
        sub_no += 1
        if m.group(3): REFS[m.group(3)] = f"{roman[sec_no-1]}-{chr(64+sub_no)}"
    elif m.group(4):
        kind = m.group(4).rstrip("*")
        if kind == "figure": fig_no += 1; REFS[m.group(5)] = str(fig_no)
        elif kind == "table": tab_no += 1; REFS[m.group(5)] = roman[tab_no - 1]
        else: alg_no += 1; REFS[m.group(5)] = str(alg_no)
    elif m.group(6):
        eq_no += 1; REFS[m.group(6)] = str(eq_no)


def set_cols(section, num):
    cols = section._sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols"); section._sectPr.append(cols)
    cols.set(qn("w:num"), str(num)); cols.set(qn("w:space"), "360")


def build(ieee: bool):
    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(10)
    st.paragraph_format.space_after = Pt(0)
    sec = doc.sections[0]
    sec.page_height = Inches(11); sec.page_width = Inches(8.5)
    for a in ("top_margin", "bottom_margin"): setattr(sec, a, Inches(0.75))
    for a in ("left_margin", "right_margin"): setattr(sec, a, Inches(0.72))
    col = {"n": 1}

    def ensure_cols(n):
        if not ieee or col["n"] == n:
            return
        s = doc.add_section(WD_SECTION.CONTINUOUS)
        s.top_margin = s.bottom_margin = Inches(0.75); s.left_margin = s.right_margin = Inches(0.72)
        set_cols(s, n); col["n"] = n

    def para(text, size=10, bold=False, italic=False, align=None, indent=True, mono=False):
        p = doc.add_paragraph()
        p.alignment = align if align is not None else WD_ALIGN_PARAGRAPH.JUSTIFY
        if indent and align is None: p.paragraph_format.first_line_indent = Inches(0.14)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
        if mono: r.font.name = "Courier New"
        return p

    # title / authors
    title = detex(re.search(r"\\title\{(.*?)\}\s*\n", TEX, re.S).group(1))
    para(title, 18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    auth = re.search(r"\\author\{(.*?)\n\}\n", TEX, re.S).group(1)
    sup = str.maketrans("0123456789", "⁰¹²³⁴⁵⁶⁷⁸⁹")
    auth = re.sub(r"\\textsuperscript\{(\d+)\}", lambda m: m.group(1).translate(sup), auth)
    for ln in re.findall(r"\\IEEEauthorblock[NA]\{(.*?)\}(?=\n|$)", auth, re.S):
        for sub in ln.split("\\\\"):
            para(detex(sub), 10.5 if "\\textit" not in sub else 10, italic="\\textit" in sub, align=WD_ALIGN_PARAGRAPH.CENTER)
    if ieee:
        ensure_cols(2)

    body = TEX[TEX.index(r"\maketitle") + len(r"\maketitle"):TEX.index(r"\end{document}")]
    pos = 0
    env_re = re.compile(r"\\begin\{(abstract|IEEEkeywords|enumerate|itemize|figure\*?|table\*?|algorithm|equation|thebibliography)\}(.*?)\\end\{\1\}", re.S)

    def emit_text(chunk):
        for blk in re.split(r"\n\s*\n", chunk):
            blk = blk.strip()
            if not blk:
                continue
            m = re.match(r"\\section\{(.*?)\}", blk)
            if m:
                ensure_cols(2)
                p = para(detex(m.group(1)).upper() if ieee else detex(m.group(1)), 10 if ieee else 11, bold=True,
                         align=WD_ALIGN_PARAGRAPH.CENTER if ieee else WD_ALIGN_PARAGRAPH.LEFT)
                p.paragraph_format.space_before = Pt(6)
                rest = blk[m.end():].strip()
                if rest: emit_text(rest)
                continue
            m = re.match(r"\\subsection\{(.*?)\}", blk)
            if m:
                ensure_cols(2)
                p = para(detex(m.group(1)), 10, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT)
                p.paragraph_format.space_before = Pt(4)
                rest = blk[m.end():].strip()
                if rest: emit_text(rest)
                continue
            ensure_cols(2)
            para(detex(blk))

    for m in env_re.finditer(body):
        emit_text(body[pos:m.start()])
        env, inner = m.group(1), m.group(2)
        pos = m.end()
        if env == "abstract":
            p = para("Abstract—" + detex(inner), 9, bold=True, indent=False)
        elif env == "IEEEkeywords":
            para("Index Terms—" + detex(inner), 9, bold=True, indent=False)
        elif env in ("enumerate", "itemize"):
            for i, it in enumerate(re.split(r"\\item\s", inner)[1:], 1):
                para(("%d) " % i if env == "enumerate" else "• ") + detex(it), indent=False)
        elif env == "equation":
            lab = re.search(r"\\label\{([^}]*)\}", inner)
            txt = detex(inner) + ("    (%s)" % REFS[lab.group(1)] if lab else "")
            para(txt, 10, align=WD_ALIGN_PARAGRAPH.CENTER)
        elif env.startswith("figure"):
            wide = env.endswith("*")
            img = re.search(r"\\includegraphics\[[^\]]*\]\{([^}]*)\}", inner).group(1)
            cap = detex(re.search(r"\\caption\{(.*)\}", inner, re.S).group(1))
            lab = re.search(r"\\label\{([^}]*)\}", inner).group(1)
            ensure_cols(1 if (wide and ieee) else col["n"])
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(ROOT / img), width=Inches(6.9 if wide else 3.3))
            para(f"Fig. {REFS[lab]}. {cap}", 8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
            if wide and ieee: ensure_cols(2)
        elif env.startswith("table"):
            wide = env.endswith("*")
            cap = detex(re.search(r"\\caption\{(.*?)\}\n", inner, re.S).group(1))
            lab = re.search(r"\\label\{([^}]*)\}", inner).group(1)
            tab = re.search(r"\\begin\{tabular\}\{[^}]*\}(.*?)\\end\{tabular\}", inner, re.S).group(1)
            rows = []
            for ln in tab.split("\\\\"):
                ln = re.sub(r"\\(toprule|midrule|bottomrule|cmidrule\([^)]*\)\{[^}]*\})", "", ln).strip()
                if not ln or ln.startswith("%"):
                    continue
                cells = [detex(c) for c in re.split(r"(?<!\\)&", ln)]
                if any(cells):
                    rows.append(cells)
            ensure_cols(1 if (wide and ieee) else col["n"])
            para(f"TABLE {REFS[lab]}. {cap}", 8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
            ncol = max(len(r) for r in rows)
            t = doc.add_table(rows=len(rows), cols=ncol); t.style = "Light Grid"
            for i, r in enumerate(rows):
                for j in range(ncol):
                    c = t.cell(i, j); c.text = r[j] if j < len(r) else ""
                    for pp in c.paragraphs:
                        for rr in pp.runs: rr.font.size = Pt(8)
            doc.add_paragraph()
            if wide and ieee: ensure_cols(2)
        elif env == "algorithm":
            cap = detex(re.search(r"\\caption\{(.*?)\}\n", inner, re.S).group(1))
            lab = re.search(r"\\label\{([^}]*)\}", inner).group(1)
            ensure_cols(2)
            para(f"Algorithm {REFS[lab]} {cap}", 9, bold=True, indent=False)
            code = re.search(r"\\begin\{algorithmic\}(?:\[\d\])?(.*?)\\end\{algorithmic\}", inner, re.S).group(1)
            depth = 0; n = 0
            for ln in code.strip().splitlines():
                ln = ln.strip()
                if not ln: continue
                if re.match(r"\\(EndFor|EndIf)", ln): depth -= 1; continue
                if ln.startswith("\\ElsIf"): depth -= 1
                n += 1
                txt = re.sub(r"\\(State|For|If|ElsIf|Return)", lambda mm: {"State": "", "For": "for ", "If": "if ", "ElsIf": "else if ", "Return": "return "}[mm.group(1)], ln)
                para(f"{n:2d}: " + "    " * max(depth, 0) + detex(txt), 8.5, indent=False, align=WD_ALIGN_PARAGRAPH.LEFT, mono=True)
                if re.match(r"\\(For|If|ElsIf)", ln): depth += 1
            doc.add_paragraph()
        elif env == "thebibliography":
            ensure_cols(2)
            para("REFERENCES" if ieee else "References", 10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER if ieee else WD_ALIGN_PARAGRAPH.LEFT)
            for i, it in enumerate(re.split(r"\\bibitem\{[^}]*\}", inner)[1:], 1):
                para(f"[{i}] " + detex(it), 8.5, indent=False, align=WD_ALIGN_PARAGRAPH.LEFT)
    emit_text(body[pos:])
    out = ROOT / ("paper_ieee.docx" if ieee else "paper.docx")
    doc.save(str(out)); print("Saved", out)


if __name__ == "__main__":
    build(False); build(True)
