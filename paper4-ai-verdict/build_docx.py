"""
Builds paper4-ai-verdict Word documents from paper.md with the same
real-number substitution (from ../ai_verdict_results.json) and image
embedding as the LaTeX path. Produces:
  paper.docx       — simple single-column Word (easy editing)
  paper_ieee.docx  — IEEE-journal-styled two-column Word (closer to Xplore)

python-docx cannot perfectly reproduce IEEEtran (true column-spanning floats).
For a pixel-perfect PDF, paper.tex on Overleaf remains canonical; these .docx
files are the editable Word approximations. Run with /usr/bin/python3.
"""
import json
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).parent
FIG = ROOT / "figures"
D = json.loads((ROOT.parent / "ai_verdict_results.json").read_text())
S = D["summary"]; T = D["tamper_detection"]; OC = D["onchain_retrievability"]
RAW = D["raw_results"]

CATEGORIES = ["reconnaissance", "privilege_escalation", "sql_injection", "reverse_shell",
              "cryptomining", "crypto_theft", "prompt_injection", "supply_chain", "unknown"]


def _cat_stats():
    by_cat = {c: [0, 0, 0] for c in CATEGORIES}
    for r in RAW:
        gt = r["ground_truth"]
        if gt not in by_cat:
            continue
        by_cat[gt][2] += 1
        if r["llm_correct"]:
            by_cat[gt][0] += 1
        if r["rule_based_correct"]:
            by_cat[gt][1] += 1
    return {c: {"n": n, "llm_pct": round(100 * lc / n, 1) if n else 0.0,
                "rule_pct": round(100 * rc / n, 1) if n else 0.0, "llm_correct": lc}
            for c, (lc, rc, n) in by_cat.items()}


CAT = _cat_stats()
mism = S["confusion_mismatches"]
recon_miscls = sum(v for k, v in mism.items() if k.endswith("->reconnaissance"))
total_errors = sum(mism.values())
llm_lat = S["llm_latency_ms"]; anchor_lat = S["anchor_latency_ms"]
unk = CAT["unknown"]

SUBS = {
    "N_PAYLOADS": f"{S['n_payloads']}",
    "RULE_ACC": f"{S['rule_based_accuracy_pct']:.1f}",
    "LLM_ACC": f"{S['llm_accuracy_pct']:.1f}",
    "ACC_GAP": f"{S['rule_based_accuracy_pct'] - S['llm_accuracy_pct']:.1f}",
    "AGREE_PCT": f"{S['llm_rule_agreement_pct']:.1f}",
    "ANCHOR_MEAN": f"{anchor_lat['mean_ms']:.2f}",
    "ANCHOR_RATE": f"{S['anchored_rate_pct']:.0f}",
    "ANCHOR_OK": f"{S['anchored_ok']}",
    "LLM_MEAN": f"{llm_lat['mean_ms']:,.1f}",
    "LLM_MEDIAN": f"{llm_lat['median_ms']:,.1f}",
    "TAMPER_TRIALS": f"{T['trials']}",
    "TAMPER_DETECTED": f"{T['detected']}",
    "TAMPER_RATE": f"{T['detection_rate_pct']:.0f}",
    "TAMPER_FP": f"{T['false_positives']}",
    "ONCHAIN_FOUND": "yes" if OC["found_on_chain"] else "no",
    "PI_LLM": f"{CAT['prompt_injection']['llm_pct']:.0f}",
    "PI_RULE": f"{CAT['prompt_injection']['rule_pct']:.0f}",
    "CT_LLM": f"{CAT['crypto_theft']['llm_pct']:.0f}",
    "CT_RULE": f"{CAT['crypto_theft']['rule_pct']:.0f}",
    "SC_LLM": f"{CAT['supply_chain']['llm_pct']:.0f}",
    "SC_RULE": f"{CAT['supply_chain']['rule_pct']:.0f}",
    "UNK_LLM_N": f"{unk['llm_correct']}",
    "UNK_N": f"{unk['n']}",
    "UNK_MISS": f"{unk['n'] - unk['llm_correct']}",
    "TOTAL_ERRORS": f"{total_errors}",
    "RECON_MISCLASS": f"{recon_miscls}",
    "RECON_MISCLASS_PCT": f"{100 * recon_miscls / total_errors:.0f}" if total_errors else "0",
}

# wide items span both columns; charts sit inline
FULLWIDTH = {
    "FIG_PIPELINE": ("diag_pipeline.png", 6.9, "Fig. 1. The dual-classification, verdict-anchoring pipeline."),
    "FIG_PER_CATEGORY": ("fig_per_category.png", 6.9, "Fig. 4. Per-category accuracy, rule-based vs. LLM classifier."),
    "TAB_CATEGORY_BODY": ("tab_category.png", 6.9, "TABLE II. Per-category accuracy and the rule-based-minus-LLM gap."),
}
INLINE_IMG = {
    "FIG_RECORD": ("diag_record.png", 3.3, "Fig. 2. The verdict record: hashed fields vs. local annotation."),
    "FIG_ACCURACY": ("fig_accuracy.png", 3.3, "Fig. 3. Headline classification accuracy."),
    "FIG_CONFUSION": ("fig_confusion.png", 3.3, "Fig. 5. Where LLM misclassifications land."),
    "FIG_LATENCY": ("fig_latency.png", 3.3, "Fig. 6. Cost of a verdict: anchoring vs. inference."),
    "FIG_TAMPER": ("fig_tamper.png", 3.3, "Fig. 7. Verdict tamper-detection outcome."),
    "TAB_RESULTS_BODY": ("tab_results.png", 3.3, "TABLE I. Summary of measured results."),
}
ALL_IMG = {**FULLWIDTH, **INLINE_IMG}


def fill(t):
    return re.sub(r"\{\{([A-Z0-9_]+)\}\}", lambda m: SUBS.get(m.group(1), m.group(0)), t)


def strip_md(t):
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t)
    t = re.sub(r"\*(.+?)\*", r"\1", t)
    t = re.sub(r"`(.+?)`", r"\1", t)
    return t


def set_cols(section, num):
    cols = section._sectPr.find(qn('w:cols'))
    if cols is None:
        cols = OxmlElement('w:cols'); section._sectPr.append(cols)
    cols.set(qn('w:num'), str(num))
    cols.set(qn('w:space'), '360')


def build(ieee: bool):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    sec = doc.sections[0]
    sec.page_height = Inches(11); sec.page_width = Inches(8.5)
    sec.top_margin = Inches(0.75); sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.72); sec.right_margin = Inches(0.72)

    md = (ROOT / "paper.md").read_text()
    lines = md.split("\n")

    # title + author block
    i = 0
    while i < len(lines) and not lines[i].startswith("# "):
        i += 1
    title = fill(lines[i][2:].strip())
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title); r.bold = True; r.font.size = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    j = i + 1
    while j < len(lines) and lines[j].strip() != "## Abstract":
        if lines[j].strip():
            ap = doc.add_paragraph(); ap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ap.add_run(strip_md(lines[j].strip())).font.size = Pt(10.5)
        j += 1
    start = j

    col_state = {"n": 1}
    if ieee:
        ns = doc.add_section(WD_SECTION.CONTINUOUS)
        ns.top_margin = Inches(0.75); ns.bottom_margin = Inches(0.75)
        ns.left_margin = Inches(0.72); ns.right_margin = Inches(0.72)
        set_cols(ns, 2); col_state["n"] = 2

    def ensure_cols(n):
        if not ieee or col_state["n"] == n:
            return
        s = doc.add_section(WD_SECTION.CONTINUOUS)
        s.top_margin = Inches(0.75); s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.72); s.right_margin = Inches(0.72)
        set_cols(s, n); col_state["n"] = n

    def add_heading(text, level):
        ensure_cols(2)
        pp = doc.add_paragraph()
        pp.paragraph_format.space_before = Pt(6); pp.paragraph_format.space_after = Pt(3)
        text = re.sub(r"^\d+\.\d*\s*", "", text)
        if level == 1:
            if ieee:
                pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rr = pp.add_run(fill(text).upper() if ieee else fill(text))
            rr.bold = True; rr.font.size = Pt(11 if not ieee else 10)
        else:
            rr = pp.add_run(fill(text)); rr.italic = True; rr.bold = True; rr.font.size = Pt(10)

    def add_para(text, justify=True, first=False):
        ensure_cols(2)
        pp = doc.add_paragraph()
        pp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
        pp.paragraph_format.first_line_indent = Inches(0 if first else 0.14)
        pp.paragraph_format.space_after = Pt(2)
        pp.add_run(strip_md(fill(text)))

    def add_img(token, wide):
        fname, w, cap = ALL_IMG[token]
        path = FIG / fname
        if not path.exists():
            return
        ensure_cols(1 if (wide and ieee) else col_state["n"])
        pp = doc.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        maxw = 6.9 if wide else 3.3
        pp.add_run().add_picture(str(path), width=Inches(min(w, maxw)))
        cc = doc.add_paragraph(); cc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cc.add_run(cap).font.size = Pt(8.5)
        if wide and ieee:
            ensure_cols(2)

    i = start
    while i < len(lines):
        s = lines[i].strip()
        m = re.fullmatch(r"\{\{([A-Z0-9_]+)\}\}", s)
        if m and m.group(1) in FULLWIDTH:
            add_img(m.group(1), wide=True); i += 1; continue
        if m and m.group(1) in INLINE_IMG:
            add_img(m.group(1), wide=False); i += 1; continue
        if s == "---" or s == "":
            i += 1; continue
        if s.startswith("## "):
            add_heading(s[3:], 1); i += 1; continue
        if s.startswith("### "):
            add_heading(s[4:], 2); i += 1; continue
        if s.startswith("# "):
            i += 1; continue
        if s.startswith("> "):
            add_para(s[2:]); i += 1; continue
        if re.match(r"^\d+\.\s", s):
            add_para(re.sub(r"^\d+\.\s", "", s)); i += 1; continue
        if s.startswith("- "):
            add_para("• " + s[2:], justify=False); i += 1; continue
        para = [s]; k = i + 1
        while k < len(lines) and lines[k].strip() and not lines[k].strip().startswith(("#", "-", ">")) \
                and not re.match(r"^\d+\.\s", lines[k].strip()) \
                and not re.fullmatch(r"\{\{[A-Z0-9_]+\}\}", lines[k].strip()):
            para.append(lines[k].strip()); k += 1
        add_para(" ".join(para), first=s.lower().startswith(("a honeypot", "index terms")))
        i = k

    out = ROOT / ("paper_ieee.docx" if ieee else "paper.docx")
    doc.save(str(out))
    leftover = set(re.findall(r"\{\{([A-Z0-9_]+)\}\}", md))
    unresolved = [t for t in leftover if t not in SUBS and t not in ALL_IMG]
    print(f"Saved {out}" + ("" if not unresolved else f"  WARNING unresolved: {sorted(unresolved)}"))


if __name__ == "__main__":
    build(ieee=False)
    build(ieee=True)
