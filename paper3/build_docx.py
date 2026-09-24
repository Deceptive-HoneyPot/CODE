"""
Builds paper3 Word documents from paper.md with the same real-number
substitution (from ../blockchain_results.json) and image embedding as the
LaTeX path. Produces:
  paper.docx       — simple single-column Word (easy editing)
  paper_ieee.docx  — IEEE-conference-styled two-column Word (closer to Xplore)

python-docx cannot perfectly reproduce IEEEtran (true column-spanning floats).
For a pixel-perfect PDF, paper.tex on Overleaf remains canonical; these .docx
files are the editable Word approximations. Run with /usr/bin/python3.
"""
import json
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).parent
FIG = ROOT / "figures"
B = json.loads((ROOT.parent / "blockchain_results.json").read_text())

al = B["anchor_latency_ms"]; bl = B["baseline_hash_latency_ms"]
vl = B["verification_latency_ms"]; gas = B["gas"]; suc = B["anchored_success"]
thr = B["throughput_rps"]; tam = B["tamper_detection"]; oc = B["onchain_retrievability"]
meta = B["meta"]
verify_ratio = round(al["mean_ms"] / vl["mean_ms"]) if vl["mean_ms"] else 0
_gv = re.search(r"v([0-9.]+)", meta["ganache_client"]); ganache_ver = _gv.group(1) if _gv else "dev"

SUBS = {
    "ANCHOR_N": f"{al['n']:,}", "ANCHOR_MEAN": f"{al['mean_ms']:.2f}",
    "ANCHOR_MEDIAN": f"{al['median_ms']:.2f}", "ANCHOR_STDEV": f"{al['stdev_ms']:.2f}",
    "ANCHOR_P95": f"{al['p95_ms']:.2f}", "ANCHOR_P99": f"{al['p99_ms']:.2f}",
    "ANCHOR_MIN": f"{al['min_ms']:.2f}", "ANCHOR_MAX": f"{al['max_ms']:.2f}",
    "BASE_N": f"{bl['n']:,}", "BASE_MEAN": f"{bl['mean_ms']:.4f}",
    "OVERHEAD_MS": f"{B['anchoring_overhead_ms']:.2f}",
    "VERIFY_MEAN": f"{vl['mean_ms']:.4f}", "VERIFY_MEDIAN": f"{vl['median_ms']:.4f}",
    "VERIFY_RATIO": f"{verify_ratio:,}",
    "GAS_MEAN": f"{gas['mean_gas']:,.0f}", "GAS_MIN": f"{gas['min_gas']:,}",
    "GAS_MAX": f"{gas['max_gas']:,}",
    "GAS_COST_ETH": f"{gas['mean_cost_eth']:.2e}",
    "ANCHOR_OK": f"{suc['anchored_ok']:,}", "ANCHOR_RATE": f"{suc['success_rate_pct']:.0f}",
    "TX_MINED": f"{B['onchain_txs_mined']:,}",
    "SEQ_RPS": f"{thr['sequential_anchored']:.2f}", "CONC_RPS": f"{thr['concurrent_anchored']:.2f}",
    "CONC_LEVEL": f"{thr['concurrency_level']}",
    "TAMPER_TRIALS": f"{tam['trials']:,}", "TAMPER_DETECTED": f"{tam['detected']:,}",
    "TAMPER_RATE": f"{tam['detection_rate_pct']:.0f}", "TAMPER_FP": f"{tam['false_positives']}",
    "ONCHAIN_FOUND": "yes" if oc["found_on_chain"] else "no",
    "CHAIN_ID": f"{meta['chain_id']}", "GANACHE_VER": ganache_ver,
    "TAB_COMPARE": "I", "TAB_RESULTS": "II",
}

# wide items span both columns; charts sit inline
FULLWIDTH = {
    "FIG_PIPELINE": ("diag_pipeline.png", 6.9, "Fig. 1. The tamper-evident anchoring pipeline."),
    "TAB_COMPARE_BODY": ("tab_compare.png", 6.9, "TABLE I. Positioning against prior blockchain-logging and blockchain-honeypot systems."),
}
INLINE_IMG = {
    "FIG_CHAIN": ("diag_threat.png", 3.3, "Fig. 2. The evidence-integrity threat model."),
    "FIG_LATENCY": ("fig_latency.png", 3.3, "Fig. 3. Anchoring cost decomposed (log scale)."),
    "FIG_GAS": ("fig_gas.png", 3.3, "Fig. 4. Measured gas per anchor vs. the base-tx floor."),
    "FIG_TAMPER": ("fig_tamper.png", 3.3, "Fig. 5. Tamper-detection outcome over all trials."),
    "TAB_RESULTS_BODY": ("tab_results.png", 3.3, "TABLE II. Summary of measured results."),
}
ALL_IMG = {**FULLWIDTH, **INLINE_IMG}


def fill(t):
    return re.sub(r"\{\{([A-Z0-9_]+)\}\}", lambda m: SUBS.get(m.group(1), m.group(0)), t)


def strip_md(t):
    t = re.sub(r"\*\*(.+?)\*\*", r"\1", t)
    t = re.sub(r"\*(.+?)\*", r"\1", t)
    t = re.sub(r"`(.+?)`", r"\1", t)
    return t


def set_cols(section, num):
    cols = section._sectPr.find(qn('w:cols'))
    if cols is None:
        cols = OxmlElement('w:cols'); section._sectPr.append(cols)
    cols.set(qn('w:num'), str(num))
    cols.set(qn('w:space'), '360')


def build(ieee: bool):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    sec = doc.sections[0]
    sec.page_height = Inches(11); sec.page_width = Inches(8.5)
    sec.top_margin = Inches(0.75); sec.bottom_margin = Inches(0.75)
    sec.left_margin = Inches(0.72); sec.right_margin = Inches(0.72)

    md = (ROOT / "paper.md").read_text()
    lines = md.split("\n")

    # title + author block
    i = 0
    while i < len(lines) and not lines[i].startswith("# "):
        i += 1
    title = fill(lines[i][2:].strip())
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(title); r.bold = True; r.font.size = Pt(18)
    p.paragraph_format.space_after = Pt(8)
    j = i + 1
    while j < len(lines) and lines[j].strip() != "## Abstract":
        if lines[j].strip():
            ap = doc.add_paragraph(); ap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            ap.add_run(strip_md(lines[j].strip())).font.size = Pt(10.5)
        j += 1
    start = j

    col_state = {"n": 1}
    if ieee:
        ns = doc.add_section(WD_SECTION.CONTINUOUS)
        ns.top_margin = Inches(0.75); ns.bottom_margin = Inches(0.75)
        ns.left_margin = Inches(0.72); ns.right_margin = Inches(0.72)
        set_cols(ns, 2); col_state["n"] = 2

    def ensure_cols(n):
        if not ieee or col_state["n"] == n:
            return
        s = doc.add_section(WD_SECTION.CONTINUOUS)
        s.top_margin = Inches(0.75); s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.72); s.right_margin = Inches(0.72)
        set_cols(s, n); col_state["n"] = n

    def add_heading(text, level):
        ensure_cols(2)
        pp = doc.add_paragraph()
        pp.paragraph_format.space_before = Pt(6); pp.paragraph_format.space_after = Pt(3)
        text = re.sub(r"^\d+\.\d*\s*", "", text)
        if level == 1:
            if ieee:
                pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rr = pp.add_run(fill(text).upper() if ieee else fill(text))
            rr.bold = True; rr.font.size = Pt(11 if not ieee else 10)
        else:
            rr = pp.add_run(fill(text)); rr.italic = True; rr.bold = True; rr.font.size = Pt(10)

    def add_para(text, justify=True, first=False):
        ensure_cols(2)
        pp = doc.add_paragraph()
        pp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if justify else WD_ALIGN_PARAGRAPH.LEFT
        pp.paragraph_format.first_line_indent = Inches(0 if first else 0.14)
        pp.paragraph_format.space_after = Pt(2)
        pp.add_run(strip_md(fill(text)))

    def add_img(token, wide):
        fname, w, cap = ALL_IMG[token]
        path = FIG / fname
        if not path.exists():
            return
        ensure_cols(1 if (wide and ieee) else col_state["n"])
        pp = doc.add_paragraph(); pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        maxw = 6.9 if wide else 3.3
        pp.add_run().add_picture(str(path), width=Inches(min(w, maxw)))
        cc = doc.add_paragraph(); cc.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cc.add_run(cap).font.size = Pt(8.5)
        if wide and ieee:
            ensure_cols(2)

    i = start
    while i < len(lines):
        s = lines[i].strip()
        m = re.fullmatch(r"\{\{([A-Z0-9_]+)\}\}", s)
        if m and m.group(1) in FULLWIDTH:
            add_img(m.group(1), wide=True); i += 1; continue
        if m and m.group(1) in INLINE_IMG:
            add_img(m.group(1), wide=False); i += 1; continue
        if s == "---" or s == "":
            i += 1; continue
        if s.startswith("## "):
            add_heading(s[3:], 1); i += 1; continue
        if s.startswith("### "):
            add_heading(s[4:], 2); i += 1; continue
        if s.startswith("# "):
            i += 1; continue
        if s.startswith("> "):
            add_para(s[2:]); i += 1; continue
        if re.match(r"^\d+\.\s", s):
            add_para(re.sub(r"^\d+\.\s", "", s)); i += 1; continue
        if s.startswith("- "):
            add_para("• " + s[2:], justify=False); i += 1; continue
        para = [s]; k = i + 1
        while k < len(lines) and lines[k].strip() and not lines[k].strip().startswith(("#", "-", ">")) \
                and not re.match(r"^\d+\.\s", lines[k].strip()) \
                and not re.fullmatch(r"\{\{[A-Z0-9_]+\}\}", lines[k].strip()):
            para.append(lines[k].strip()); k += 1
        add_para(" ".join(para), first=s.lower().startswith(("a honeypot", "index terms")))
        i = k

    out = ROOT / ("paper_ieee.docx" if ieee else "paper.docx")
    doc.save(str(out))
    leftover = set(re.findall(r"\{\{([A-Z0-9_]+)\}\}", md))
    unresolved = [t for t in leftover if t not in SUBS and t not in ALL_IMG]
    print(f"Saved {out}" + ("" if not unresolved else f"  WARNING unresolved: {sorted(unresolved)}"))


if __name__ == "__main__":
    build(ieee=False)
    build(ieee=True)
